"""
Script Generator Presentasi PowerPoint (.pptx)
Tugas Kuliah: Analisis dan Pengujian Sistem
Topik: Root Cause Analysis (RCA) & Defect Identification pada Mini E-Commerce Order System
Mahasiswa: Zainul Mutawakkil (NIM: 23EO10003)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ================= PALET WARNA PROFESIONAL =================
COLOR_PRIMARY_DARK  = RGBColor(15, 23, 42)     # Slate 900 (Deep Navy Dark)
COLOR_SECONDARY_DARK= RGBColor(30, 41, 59)     # Slate 800
COLOR_BG_LIGHT      = RGBColor(248, 250, 252)  # Slate 50 (Very light gray-blue)
COLOR_CARD_BG       = RGBColor(255, 255, 255)  # Pure White
COLOR_CARD_BORDER   = RGBColor(226, 232, 240)  # Slate 200
COLOR_ACCENT_BLUE   = RGBColor(37, 99, 235)    # Royal Blue
COLOR_ACCENT_CYAN   = RGBColor(14, 165, 233)   # Sky Blue
COLOR_DANGER_RED    = RGBColor(220, 38, 38)    # Crimson Red
COLOR_SUCCESS_GREEN = RGBColor(16, 185, 129)   # Emerald Green
COLOR_WARNING_AMBER = RGBColor(245, 158, 11)   # Amber
COLOR_TEXT_MAIN     = RGBColor(15, 23, 42)     # Deep Charcoal
COLOR_TEXT_MUTED    = RGBColor(100, 116, 139)  # Slate Muted
COLOR_TEXT_WHITE    = RGBColor(255, 255, 255)  # White

FONT_FAMILY = "Segoe UI"
FONT_CODE = "Consolas"

def create_deck():
    prs = Presentation()
    # Format Widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, category: str, title: str, slide_num: int):
        # Header bar background
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Category pill/subtitle
        p0 = tf.paragraphs[0]
        p0.text = category.upper()
        p0.font.name = FONT_FAMILY
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_ACCENT_BLUE
        p0.space_after = Pt(4)

        # Title
        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_MAIN

        # Slide Number in corner
        num_box = slide.shapes.add_textbox(Inches(11.5), Inches(0.4), Inches(1.0), Inches(0.5))
        ntf = num_box.text_frame
        np = ntf.paragraphs[0]
        np.text = f"{slide_num:02d} / 16"
        np.font.name = FONT_FAMILY
        np.font.size = Pt(12)
        np.font.color.rgb = COLOR_TEXT_MUTED
        np.alignment = PP_ALIGN.RIGHT

    def add_card(slide, left, top, width, height, title="", title_color=COLOR_TEXT_MAIN, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
        # Background rectangle
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()

        # Text container inside card
        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        if title:
            p = tf.paragraphs[0]
            p.text = title
            p.font.name = FONT_FAMILY
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = title_color
            p.space_after = Pt(8)
            return tf, True
        return tf, False

    # =========================================================================
    # SLIDE 1: COVER SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, COLOR_PRIMARY_DARK)

    # Accent decorative bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.5), Inches(0.15), Inches(4.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_ACCENT_CYAN
    bar.line.fill.background()

    # Title Box
    title_box = s1.shapes.add_textbox(Inches(1.4), Inches(1.5), Inches(10.5), Inches(4.5))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "TUGAS BESAR ANALISIS DAN PENGUJIAN SISTEM"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_CYAN
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "ROOT CAUSE ANALYSIS (RCA)\n& DEFECT IDENTIFICATION"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_WHITE
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "Studi Kasus Penemuan 3 Defect Kritis, Analisis 5-Whys, Diagram Fishbone, serta Verifikasi Perbaikan pada Mini E-Commerce Order System"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(15)
    p.font.color.rgb = RGBColor(203, 213, 225)
    p.space_after = Pt(32)

    # Author card inside cover
    p = tf1.add_paragraph()
    p.text = "Disusun Oleh : Zainul Mutawakkil  |  NIM : 23EO10003"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_WHITE

    p = tf1.add_paragraph()
    p.text = "Mata Kuliah : Analisis dan Pengujian Sistem  •  Tahun Akademik 2026"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(148, 163, 184)

    # =========================================================================
    # SLIDE 2: GAMBARAN SISTEM MINI PROGRAM
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, COLOR_BG_LIGHT)
    add_header(s2, "Arsitektur Perangkat Lunak", "Gambaran Umum Mini E-Commerce Order System", 2)

    # 4 Cards for modules
    modules = [
        ("Inventory Manager (inventory.py)", "Mengelola ketersediaan stok fisik barang di gudang. Bertanggung jawab atas operasi cek stok dan pengurangan stok saat checkout.", COLOR_ACCENT_BLUE),
        ("Pricing Engine (pricing.py)", "Menghitung subtotal belanja, diskon bertingkat (persentase & voucher nominal), serta tarif pajak pertambahan nilai (PPN 11%).", COLOR_ACCENT_BLUE),
        ("Payment Gateway (payment.py)", "Simulasi integrasi gateway pembayaran pihak ketiga (bank/fintech) dengan skenario timeout jaringan dan saldo tidak cukup.", COLOR_ACCENT_BLUE),
        ("Order Service (order_service.py)", "Orchestrator utama yang menghubungkan pemotongan stok, kalkulasi tagihan, dan pemanggilan payment gateway.", COLOR_ACCENT_BLUE),
    ]

    for i, (m_title, m_desc, col) in enumerate(modules):
        col_idx = i % 2
        row_idx = i // 2
        left = Inches(0.8 + col_idx * 5.9)
        top = Inches(1.8 + row_idx * 2.5)
        tf, _ = add_card(s2, left, top, Inches(5.6), Inches(2.2), f"📦 Modul: {m_title}", col)
        
        p = tf.add_paragraph()
        p.text = m_desc
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_MAIN

    # Bottom summary banner
    tf_sum, _ = add_card(s2, Inches(0.8), Inches(6.1), Inches(11.5), Inches(0.9), "", border_color=COLOR_ACCENT_BLUE, bg_color=RGBColor(239, 246, 255))
    p = tf_sum.paragraphs[0]
    p.text = "🎯 Tujuan Mini Program: Menyediakan repositori modular realistis untuk mendemonstrasikan bagaimana celah desain kecil pada konkurensi, tipe data, dan alur transaksi dapat berkembang menjadi kegagalan sistemik berisiko tinggi."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_BLUE

    # =========================================================================
    # SLIDE 3: LANDASAN TEORI - TAKSONOMI KEGAGALAN & RCA
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, COLOR_BG_LIGHT)
    add_header(s3, "Landasan Teori & Konsep", "Taksonomi Kegagalan Software & Metodologi RCA", 3)

    # Left: Fault -> Error -> Failure
    tf_tax, _ = add_card(s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "1. Taksonomi Kegagalan (Standar IEEE 610.12)", COLOR_ACCENT_BLUE)
    
    stages = [
        ("HUMAN MISTAKE (Kekeliruan Manusia)", "Kekeliruan kognitif atau asumsi salah oleh programmer saat merancang/menulis kode.", COLOR_TEXT_MUTED),
        ("FAULT / DEFECT (Cacat pada Kode)", "Cacat statis di dalam baris kode program (misal ketiadaan lock, penggunaan float).", COLOR_DANGER_RED),
        ("ERROR (Kondisi State Internal Korup)", "Penyimpangan nilai internal di memori saat kode dieksekusi (stok bernilai -9, saldo minus).", COLOR_WARNING_AMBER),
        ("FAILURE (Kegagalan Eksternal)", "Kegagalan fungsional yang teramati langsung oleh user atau sistem luar (oversell, komplain).", COLOR_DANGER_RED)
    ]
    for s_title, s_desc, s_col in stages:
        p = tf_tax.add_paragraph()
        p.text = f"• {s_title}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = s_col
        p.space_after = Pt(2)

        p2 = tf_tax.add_paragraph()
        p2.text = f"  {s_desc}"
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_MAIN
        p2.space_after = Pt(8)

    # Right: Metodologi RCA (5-Whys & Fishbone)
    tf_rca, _ = add_card(s3, Inches(6.7), Inches(1.8), Inches(5.6), Inches(5.1), "2. Metodologi Root Cause Analysis (RCA)", COLOR_ACCENT_BLUE)
    
    methods = [
        ("Metode 5-Whys (Sakichi Toyoda)", "Teknik investigasi iteratif dengan menanyakan 'Mengapa' minimal 5 kali secara beruntun dari gejala luar hingga menembus akar masalah terdalam (organisasi, desain arsitektur, atau proses QA)."),
        ("Ishikawa / Fishbone Diagram (Kaoru Ishikawa)", "Diagram sebab-akibat untuk memetakan akar penyebab ke dalam 4 kategori komprehensif:\n - Man (Keahlian & kelalaian manusia)\n - Machine (Infrastruktur, thread, hardware)\n - Method (Metodologi pengujian, code review)\n - Material (Spesifikasi, batasan bisnis)."),
        ("Corrective & Preventive Actions (CAPA)", "Output terpenting RCA: Solusi kuratif langsung (bug fix) dan solusi preventif jangka panjang (otomasi QA) agar insiden tidak terulang.")
    ]
    for m_title, m_desc in methods:
        p = tf_rca.add_paragraph()
        p.text = f"• {m_title}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT_BLUE
        p.space_after = Pt(2)

        p2 = tf_rca.add_paragraph()
        p2.text = f"  {m_desc}"
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_MAIN
        p2.space_after = Pt(10)

    # =========================================================================
    # SLIDE 4: RINGKASAN 3 DEFECT / FAULT
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, COLOR_BG_LIGHT)
    add_header(s4, "Hasil Identifikasi", "Ringkasan 3 Defect/Fault Kritis yang Ditemukan", 4)

    defects_summary = [
        ("FAULT 01", "Concurrency Race Condition (Inventory Oversell)", "Modul: inventory.py\nKategori: Thread-Safety & Concurrency\nSeverity: CRITICAL / HIGH", "Stok fisik tersisa 1 unit diserbu 10 user konkuren. Seluruh 10 pesanan lolos dan stok menjadi -9 (minus).", "Pencegahan overselling dengan Reentrant Mutex Lock (threading.RLock).", COLOR_DANGER_RED),
        ("FAULT 02", "Float Precision Leak & Unbounded Discount", "Modul: pricing.py\nKategori: Business Logic & Precision\nSeverity: HIGH (Financial Loss)", "Penggunaan float menimbulkan biner artefak (3*0.1=0.30000000000000004) dan voucher tumpuk membuat tagihan -Rp 22.200.", "Migrasi ke decimal.Decimal dan validasi batas atas min(discount, subtotal).", COLOR_WARNING_AMBER),
        ("FAULT 03", "Phantom Stock Deduction (No Transaction Rollback)", "Modul: order_service.py\nKategori: Data Integrity & Transactions\nSeverity: CRITICAL / DATA LOSS", "Payment gateway mengalami network timeout. Pesanan gagal tetapi stok gudang hilang permanen tanpa kembali.", "Penerapan Compensating Transaction (try-except-rollback) untuk menjamin atomisitas.", COLOR_DANGER_RED),
    ]

    for i, (badge, d_name, d_meta, d_symptom, d_fix, col) in enumerate(defects_summary):
        left = Inches(0.8 + i * 3.9)
        top = Inches(1.8)
        tf, _ = add_card(s4, left, top, Inches(3.7), Inches(5.1), f"⚠️ {badge}", col)

        p = tf.add_paragraph()
        p.text = d_name
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_MAIN
        p.space_after = Pt(8)

        p = tf.add_paragraph()
        p.text = d_meta
        p.font.name = FONT_CODE
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(10)

        p = tf.add_paragraph()
        p.text = "💥 Gejala Kegagalan:"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_DANGER_RED

        p = tf.add_paragraph()
        p.text = d_symptom
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_MAIN
        p.space_after = Pt(10)

        p = tf.add_paragraph()
        p.text = "🛡️ Solusi Perbaikan:"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_SUCCESS_GREEN

        p = tf.add_paragraph()
        p.text = d_fix
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 5: DEEP DIVE FAULT 1
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, COLOR_BG_LIGHT)
    add_header(s5, "Studi Kasus Defect 1", "Fault 1: Concurrency Race Condition pada Stok", 5)

    tf_f1_left, _ = add_card(s5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "Gejala & Bukti Pengujian Empiris", COLOR_DANGER_RED)
    p = tf_f1_left.add_paragraph()
    p.text = "Skenario Pengujian (test_faults_reproducer.py):"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.bold = True
    p.space_after = Pt(4)

    bullet_f1 = [
        "Stok awal produk 'Laptop Gaming': 1 unit.",
        "10 Thread pelanggan melepaskan checkout secara simultan (Barrier Synchronization).",
        "Ekspektasi: Tepat 1 transaksi berhasil, 9 transaksi ditolak.",
        "Hasil Aktual: 10 transaksi berhasil disetujui!",
        "Stok Akhir di Gudang: -9 (Minus Sembilan)."
    ]
    for b in bullet_f1:
        p = tf_f1_left.add_paragraph()
        p.text = f"• {b}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.space_after = Pt(4)

    p = tf_f1_left.add_paragraph()
    p.text = "\nLog Hasil Test Runner:\n>>> STOK AWAL FISIK     : 1\n>>> TRANSAKSI SUKSES    : 10 (Seharusnya maks 1!)\n>>> STOK AKHIR DI GUDANG: -9\n>>> FAULT 1 BERHASIL DIBUKTIKAN (OVERSELLING)"
    p.font.name = FONT_CODE
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_DANGER_RED

    tf_f1_right, _ = add_card(s5, Inches(6.7), Inches(1.8), Inches(5.6), Inches(5.1), "Analisis Cacat Kode (Source Code Flaw)", COLOR_TEXT_MAIN)
    p = tf_f1_right.add_paragraph()
    p.text = "File: src/shop/inventory.py (VulnerableInventoryManager)"
    p.font.name = FONT_CODE
    p.font.size = Pt(11)
    p.font.bold = True
    p.space_after = Pt(8)

    code_vulnerable_1 = (
        "def deduct_stock(self, sku: str, quantity: int):\n"
        "    # 1. Pengecekan stok (Time of Check)\n"
        "    if self._products[sku].stock >= quantity:\n"
        "        # 2. Jeda waktu / Context switch thread\n"
        "        time.sleep(0.005)\n"
        "        # 3. Pengurangan stok (Time of Use) - CACAT!\n"
        "        self._products[sku].stock -= quantity\n"
        "        return True\n"
        "    return False"
    )
    p = tf_f1_right.add_paragraph()
    p.text = code_vulnerable_1
    p.font.name = FONT_CODE
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_DANGER_RED
    p.space_after = Pt(10)

    p = tf_f1_right.add_paragraph()
    p.text = "🔴 Titik Lemah (TOCTOU):\nDi antara baris pengecekan 'if' dan eksekusi pengurangan '-=', thread lain sempat membaca nilai stok yang belum berubah. Tidak ada mutual exclusion (lock) yang melindungi critical section."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 6: RCA FAULT 1 (5-WHYS & FISHBONE)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, COLOR_BG_LIGHT)
    add_header(s6, "Root Cause Analysis", "RCA Defect 1: 5-Whys & Fishbone Diagram", 6)

    tf_5w_1, _ = add_card(s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "5-Whys Analysis (Defect 1)", COLOR_ACCENT_BLUE)
    whys_1 = [
        ("Why 1", "Mengapa stok bisa bernilai -9?", "Karena ada 10 transaksi yang disetujui bersamaan pada stok 1."),
        ("Why 2", "Mengapa 10 transaksi bisa disetujui?", "Karena seluruh 10 thread mendeteksi stok >= 1 masih terpenuhi."),
        ("Why 3", "Mengapa semua thread membaca nilai yang sama?", "Pembacaan dilakukan paralel sebelum thread lain sempat mengurangi stok."),
        ("Why 4", "Mengapa operasi baca-kurang tidak sinkron?", "Tidak ada implementasi Mutex Lock pada critical section data stok."),
        ("Why 5", "AKAR MASALAH (Root Cause):", "Developer mengira operasi memori otomatis atomik & pipeline QA tidak memiliki concurrency stress test.")
    ]
    for w_num, w_q, w_a in whys_1:
        p = tf_5w_1.add_paragraph()
        p.text = f"• {w_num}: {w_q}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT_BLUE if "AKAR" not in w_num else COLOR_DANGER_RED
        
        p2 = tf_5w_1.add_paragraph()
        p2.text = f"   👉 {w_a}"
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = COLOR_TEXT_MAIN
        p2.space_after = Pt(6)

    tf_fb_1, _ = add_card(s6, Inches(6.7), Inches(1.8), Inches(5.6), Inches(5.1), "Ishikawa (Fishbone) 4M Category", COLOR_ACCENT_BLUE)
    fishbone_1 = [
        ("MAN (SDM & Pengetahuan)", "Kurangnya pemahaman developer terkait Thread Safety & Race Condition pada arsitektur web concurrently-driven."),
        ("MACHINE (Infrastruktur / Runtime)", "Runtime Python multi-threading melakukan preemptive thread switching saat operasi I/O atau sleep."),
        ("METHOD (Prosedur & Pengujian)", "Pengujian hanya berfokus pada single-threaded unit test. Tidak ada checklist Concurrency Code Review."),
        ("MATERIAL (Spesifikasi & Desain)", "Spesifikasi SRS tidak mendefinisikan batasan load transaksi bersamaan pada event Flash Sale.")
    ]
    for cat, desc in fishbone_1:
        p = tf_fb_1.add_paragraph()
        p.text = f"🐟 {cat}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT_BLUE
        
        p2 = tf_fb_1.add_paragraph()
        p2.text = f"   {desc}"
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = COLOR_TEXT_MAIN
        p2.space_after = Pt(8)

    # =========================================================================
    # SLIDE 7: SOLUSI & VERIFIKASI FAULT 1
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, COLOR_BG_LIGHT)
    add_header(s7, "Solusi & Verifikasi", "Perbaikan Defect 1: Thread-Safe Mutex Lock", 7)

    tf_sol_1, _ = add_card(s7, Inches(0.8), Inches(1.8), Inches(6.0), Inches(5.1), "Implementasi Kode Perbaikan (Fixed)", COLOR_SUCCESS_GREEN)
    code_fixed_1 = (
        "class FixedInventoryManager:\n"
        "    def __init__(self):\n"
        "        self._products = {}\n"
        "        # KUNCI MUTEX REENTRANT\n"
        "        self._lock = threading.RLock()\n\n"
        "    def deduct_stock(self, sku: str, quantity: int):\n"
        "        with self._lock:  # Critical Section Terkunci\n"
        "            if self._products[sku].stock >= quantity:\n"
        "                time.sleep(0.005)\n"
        "                self._products[sku].stock -= quantity\n"
        "                return True\n"
        "            return False"
    )
    p = tf_sol_1.add_paragraph()
    p.text = code_fixed_1
    p.font.name = FONT_CODE
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_SUCCESS_GREEN
    p.space_after = Pt(10)

    p = tf_sol_1.add_paragraph()
    p.text = "💡 Prinsip Solusi:\nBlok 'with self._lock' memastikan operasi pengecekan dan pengurangan stok bersifat ATOMIK. Hanya 1 thread yang diizinkan mengakses data stok dalam satu waktu."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_MAIN

    tf_ver_1, _ = add_card(s7, Inches(7.1), Inches(1.8), Inches(5.2), Inches(5.1), "Hasil Verifikasi Otomatis (test_faults_fixed.py)", COLOR_ACCENT_BLUE)
    p = tf_ver_1.add_paragraph()
    p.text = "Hasil Eksekusi Unit Test Verifikasi:\n"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.bold = True

    verif_items_1 = [
        ("Stok Awal Fisik", "1 Unit"),
        ("Permintaan Masuk", "10 Thread Paralel"),
        ("Transaksi Sukses", "1 Transaksi (100% Sesuai)"),
        ("Transaksi Ditolak", "9 Transaksi (Aman ditolak)"),
        ("Stok Akhir Gudang", "0 Unit (Tidak Pernah Minus!)"),
        ("Status Pengujian", "PASSED / OK (Zero Overselling)")
    ]
    for k, v in verif_items_1:
        p = tf_ver_1.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True if "Status" in k or "Akhir" in k else False
        p.font.color.rgb = COLOR_SUCCESS_GREEN if "Status" in k or "Akhir" in k else COLOR_TEXT_MAIN
        p.space_after = Pt(6)

    # =========================================================================
    # SLIDE 8: DEEP DIVE FAULT 2
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, COLOR_BG_LIGHT)
    add_header(s8, "Studi Kasus Defect 2", "Fault 2: Float Precision Leak & Unbounded Discount", 8)

    tf_f2_left, _ = add_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "Gejala & Bukti Pengujian Empiris", COLOR_DANGER_RED)
    p = tf_f2_left.add_paragraph()
    p.text = "Skenario Kasus Nyata:"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.bold = True
    p.space_after = Pt(6)

    bullet_f2 = [
        "Kasus A: Pecahan Moneter (3 x Rp 0.1)\n  Hasil float: Rp 0.30000000000000004\n  Dampak: Residu pembulatan merusak laporan keuangan audit.",
        "Kasus B: Promo Tumpuk Tanpa Batas\n  Harga Buku: Rp 100.000\n  Diskon Kupon 50%: Rp 50.000\n  Voucher Tambahan: Rp 70.000\n  Total Diskon: Rp 120.000 (Melebihi Subtotal!)\n  Dasar Pajak: Rp -20.000\n  Pajak (PPN 11%): Rp -2.200\n  TOTAL TAGIHAN: -Rp 22.200 (NEGATIF!)"
    ]
    for b in bullet_f2:
        p = tf_f2_left.add_paragraph()
        p.text = f"• {b}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.space_after = Pt(8)

    tf_f2_right, _ = add_card(s8, Inches(6.7), Inches(1.8), Inches(5.6), Inches(5.1), "Analisis Cacat Kode (Source Code Flaw)", COLOR_TEXT_MAIN)
    p = tf_f2_right.add_paragraph()
    p.text = "File: src/shop/pricing.py (VulnerablePricingEngine)"
    p.font.name = FONT_CODE
    p.font.size = Pt(11)
    p.font.bold = True
    p.space_after = Pt(8)

    code_vulnerable_2 = (
        "def calculate_totals(self, items, percentage_discount, fixed_discount):\n"
        "    subtotal = sum(i.quantity * i.unit_price for i in items)\n"
        "    # Float inexactness\n"
        "    discount_pct = subtotal * (percentage_discount / 100.0)\n"
        "    # Akumulasi tanpa batas atas (No Ceiling Cap)\n"
        "    total_discount = discount_pct + fixed_discount\n"
        "    taxable_amount = subtotal - total_discount # BISA MINUS\n"
        "    tax = taxable_amount * self.tax_rate\n"
        "    return {'total': taxable_amount + tax}"
    )
    p = tf_f2_right.add_paragraph()
    p.text = code_vulnerable_2
    p.font.name = FONT_CODE
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_DANGER_RED
    p.space_after = Pt(10)

    p = tf_f2_right.add_paragraph()
    p.text = "🔴 Titik Lemah:\n1. Tipe data 'float' IEEE 754 tidak tepat untuk kalkulasi moneter.\n2. Tidak ada boundary check bahwa total_discount tidak boleh melebihi subtotal, dan total tagihan tidak boleh < 0."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 9: RCA FAULT 2 (5-WHYS & FISHBONE)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, COLOR_BG_LIGHT)
    add_header(s9, "Root Cause Analysis", "RCA Defect 2: 5-Whys & Fishbone Diagram", 9)

    tf_5w_2, _ = add_card(s9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "5-Whys Analysis (Defect 2)", COLOR_ACCENT_BLUE)
    whys_2 = [
        ("Why 1", "Mengapa total invoice bernilai negatif (-Rp 22.200)?", "Karena nilai potongan diskon lebih besar dari harga pokok barang."),
        ("Why 2", "Mengapa diskon bisa lebih besar dari subtotal?", "Sistem menjumlahkan voucher persentase dan nominal tanpa batasan."),
        ("Why 3", "Mengapa tidak dibatasi batas maksimal diskon?", "Developer menerapkan formula matematika langsung tanpa aturan bisnis 'cap'."),
        ("Why 4", "Mengapa ada angka desimal biner aneh (0.3000...4)?", "Menggunakan tipe primitif float bawaan CPU daripada representasi desimal moneter."),
        ("Why 5", "AKAR MASALAH (Root Cause):", "Ketiadaan Value Object 'Money' berbasis Decimal dan tidak dilakukannya Boundary Value Analysis (BVA) pada skenario diskon.")
    ]
    for w_num, w_q, w_a in whys_2:
        p = tf_5w_2.add_paragraph()
        p.text = f"• {w_num}: {w_q}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT_BLUE if "AKAR" not in w_num else COLOR_DANGER_RED
        
        p2 = tf_5w_2.add_paragraph()
        p2.text = f"   👉 {w_a}"
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = COLOR_TEXT_MAIN
        p2.space_after = Pt(6)

    tf_fb_2, _ = add_card(s9, Inches(6.7), Inches(1.8), Inches(5.6), Inches(5.1), "Ishikawa (Fishbone) 4M Category", COLOR_ACCENT_BLUE)
    fishbone_2 = [
        ("MAN (SDM & Pengetahuan)", "Developer belum memahami keterbatasan representasi biner IEEE 754 float untuk perhitungan akuntansi."),
        ("MACHINE (Hardware & Floating Point)", "ALU komputer merepresentasikan bilangan pecahan dalam basis 2 biner, bukan basis 10 desimal."),
        ("METHOD (Pengujian & Verifikasi)", "Unit testing hanya menguji angka bulat sederhana (happy path); tidak menguji nilai batas ekstrim (BVA)."),
        ("MATERIAL (Aturan Bisnis & Schema)", "Spesifikasi promosi dari tim marketing tidak mendefinisikan aturan kombinasi voucher (stackable rule).")
    ]
    for cat, desc in fishbone_2:
        p = tf_fb_2.add_paragraph()
        p.text = f"🐟 {cat}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT_BLUE
        
        p2 = tf_fb_2.add_paragraph()
        p2.text = f"   {desc}"
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = COLOR_TEXT_MAIN
        p2.space_after = Pt(8)

    # =========================================================================
    # SLIDE 10: SOLUSI & VERIFIKASI FAULT 2
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, COLOR_BG_LIGHT)
    add_header(s10, "Solusi & Verifikasi", "Perbaikan Defect 2: Decimal Moneter & Boundary Cap", 10)

    tf_sol_2, _ = add_card(s10, Inches(0.8), Inches(1.8), Inches(6.0), Inches(5.1), "Implementasi Kode Perbaikan (Fixed)", COLOR_SUCCESS_GREEN)
    code_fixed_2 = (
        "from decimal import Decimal, ROUND_HALF_UP\n\n"
        "class FixedPricingEngine:\n"
        "    def calculate_totals(self, items, pct_disc, fixed_disc):\n"
        "        subtotal = sum(Decimal(str(i.qty))*Decimal(str(i.price)) ...)\n"
        "        safe_pct = max(0.0, min(100.0, pct_disc))\n"
        "        disc_pct = subtotal * (Decimal(str(safe_pct)) / Decimal('100'))\n"
        "        disc_fixed = max(Decimal('0'), Decimal(str(fixed_disc)))\n"
        "        # BOUNDARY CEILING CAP (Diskon maks = subtotal)\n"
        "        effective_discount = min(disc_pct + disc_fixed, subtotal)\n"
        "        taxable = max(Decimal('0'), subtotal - effective_discount)\n"
        "        tax = (taxable * self.tax_rate).quantize(Decimal('0.01'))\n"
        "        total = (taxable + tax).quantize(Decimal('0.01'))\n"
        "        return {'subtotal': subtotal, 'discount': effective_discount, 'total': total}"
    )
    p = tf_sol_2.add_paragraph()
    p.text = code_fixed_2
    p.font.name = FONT_CODE
    p.font.size = Pt(9.5)
    p.font.color.rgb = COLOR_SUCCESS_GREEN
    p.space_after = Pt(8)

    p = tf_sol_2.add_paragraph()
    p.text = "💡 Prinsip Solusi:\n1. Menggunakan Decimal untuk presisi matematis eksak.\n2. Fungsi min() menjamin diskon tidak pernah melebihi harga belanja.\n3. Nilai total diproteksi minimal Rp 0.00."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_MAIN

    tf_ver_2, _ = add_card(s10, Inches(7.1), Inches(1.8), Inches(5.2), Inches(5.1), "Hasil Verifikasi Otomatis (test_faults_fixed.py)", COLOR_ACCENT_BLUE)
    p = tf_ver_2.add_paragraph()
    p.text = "Hasil Eksekusi Unit Test Verifikasi:\n"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.bold = True

    verif_items_2 = [
        ("Subtotal Belanja", "Rp 100,000.00"),
        ("Diskon Terpasang", "50% + Rp 70,000 (Total Rp 120,000)"),
        ("Diskon Efektif", "Rp 100,000.00 (Dibatasi setara subtotal)"),
        ("Dasar Pajak & PPN", "Rp 0.00 (Tidak pernah minus!)"),
        ("TOTAL TAGIHAN AKHIR", "Rp 0.00 (Aman, zero leakage)"),
        ("Uji Presisi 3 * 0.1", "Tepat 0.30 (Tanpa biner artefak)"),
        ("Status Pengujian", "PASSED / OK (Presisi Terjamin)")
    ]
    for k, v in verif_items_2:
        p = tf_ver_2.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True if "Status" in k or "TOTAL" in k else False
        p.font.color.rgb = COLOR_SUCCESS_GREEN if "Status" in k or "TOTAL" in k else COLOR_TEXT_MAIN
        p.space_after = Pt(5)

    # =========================================================================
    # SLIDE 11: DEEP DIVE FAULT 3
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, COLOR_BG_LIGHT)
    add_header(s11, "Studi Kasus Defect 3", "Fault 3: Phantom Stock Deduction (No Rollback)", 11)

    tf_f3_left, _ = add_card(s11, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "Gejala & Bukti Pengujian Empiris", COLOR_DANGER_RED)
    p = tf_f3_left.add_paragraph()
    p.text = "Skenario Kasus Nyata:"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.bold = True
    p.space_after = Pt(6)

    bullet_f3 = [
        "Stok awal produk 'Smartphone Flagship': 5 unit.",
        "Pelanggan memesan 2 unit.",
        "Saat proses charge ke payment gateway bank, terjadi Network Timeout (HTTP 504).",
        "Sistem menangkap exception dan menandai status order: FAILED.",
        "Pelanggan tidak dikenakan biaya sama sekali.",
        "ANOMALI PARAH:\nStok di gudang berkurang menjadi 3 unit dan TIDAK DIKEMBALIKAN!",
        "Dampak: 2 Unit barang hilang secara misterius (Phantom Deduction)."
    ]
    for b in bullet_f3:
        p = tf_f3_left.add_paragraph()
        p.text = f"• {b}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.space_after = Pt(4)

    tf_f3_right, _ = add_card(s11, Inches(6.7), Inches(1.8), Inches(5.6), Inches(5.1), "Analisis Cacat Kode (Source Code Flaw)", COLOR_TEXT_MAIN)
    p = tf_f3_right.add_paragraph()
    p.text = "File: src/shop/order_service.py (VulnerableOrderService)"
    p.font.name = FONT_CODE
    p.font.size = Pt(11)
    p.font.bold = True
    p.space_after = Pt(8)

    code_vulnerable_3 = (
        "def checkout(self, customer, items, ...):\n"
        "    # 1. Deduksi stok dilakukan di awal\n"
        "    for item in items:\n"
        "        self.inventory.deduct_stock(item.sku, item.quantity)\n\n"
        "    try:\n"
        "        # 2. Pembayaran ke gateway pihak ketiga\n"
        "        self.gateway.process_charge(...)\n"
        "        return order\n"
        "    except PaymentError as err:\n"
        "        # 3. CACAT: Error ditangkap tapi LUPA ROLLBACK STOK!\n"
        "        order.status = OrderStatus.FAILED\n"
        "        return order  # Stok gudang dibiarkan hilang!"
    )
    p = tf_f3_right.add_paragraph()
    p.text = code_vulnerable_3
    p.font.name = FONT_CODE
    p.font.size = Pt(9.5)
    p.font.color.rgb = COLOR_DANGER_RED
    p.space_after = Pt(10)

    p = tf_f3_right.add_paragraph()
    p.text = "🔴 Titik Lemah:\nKetiadaan prinsip atomisitas transaksi (All-or-Nothing). Alur multi-langkah tidak memiliki blok kompensasi pemulihan state saat terjadi kegagalan parsial."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 12: RCA FAULT 3 (5-WHYS & FISHBONE)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12, COLOR_BG_LIGHT)
    add_header(s12, "Root Cause Analysis", "RCA Defect 3: 5-Whys & Fishbone Diagram", 12)

    tf_5w_3, _ = add_card(s12, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "5-Whys Analysis (Defect 3)", COLOR_ACCENT_BLUE)
    whys_3 = [
        ("Why 1", "Mengapa stok gudang berkurang padahal pesanan batal?", "Karena pemotongan stok dilakukan sebelum otorisasi gateway berhasil."),
        ("Why 2", "Mengapa stok tidak dikembalikan saat payment error?", "Blok catch exception hanya mencatat pesan error tanpa memanggil restore_stock."),
        ("Why 3", "Mengapa tidak ada mekanisme kompensasi rollback?", "Layanan checkout tidak dirancang mengikuti pola transaksi atomik."),
        ("Why 4", "Mengapa masalah ini tidak ditemukan saat fase testing?", "Pengujian tim QA hanya menguji skenario sukses (Happy Path testing)."),
        ("Why 5", "AKAR MASALAH (Root Cause):", "Ketiadaan pola Compensating Transaction (Saga Pattern) dan tidak diterapkannya Negative & Fault-Injection Testing.")
    ]
    for w_num, w_q, w_a in whys_3:
        p = tf_5w_3.add_paragraph()
        p.text = f"• {w_num}: {w_q}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT_BLUE if "AKAR" not in w_num else COLOR_DANGER_RED
        
        p2 = tf_5w_3.add_paragraph()
        p2.text = f"   👉 {w_a}"
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = COLOR_TEXT_MAIN
        p2.space_after = Pt(6)

    tf_fb_3, _ = add_card(s12, Inches(6.7), Inches(1.8), Inches(5.6), Inches(5.1), "Ishikawa (Fishbone) 4M Category", COLOR_ACCENT_BLUE)
    fishbone_3 = [
        ("MAN (SDM & Pola Pikir)", "Developer berfokus pada alur normal dan mengabaikan penanganan kondisi abnormal / parsial error."),
        ("MACHINE (Jaringan & Provider Eksternal)", "Ketergantungan pada 3rd-party Payment Gateway yang rentan fluktuasi koneksi dan timeout."),
        ("METHOD (Pengujian & Error Handling)", "Ketiadaan pengujian Chaos Engineering / Fault-Injection untuk mensimulasikan kegagalan jaringan."),
        ("MATERIAL (Arsitektur Desain)", "Arsitektur tidak memiliki Distributed Transaction Coordinator atau job rekonsiliasi stok otomatis.")
    ]
    for cat, desc in fishbone_3:
        p = tf_fb_3.add_paragraph()
        p.text = f"🐟 {cat}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT_BLUE
        
        p2 = tf_fb_3.add_paragraph()
        p2.text = f"   {desc}"
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = COLOR_TEXT_MAIN
        p2.space_after = Pt(8)

    # =========================================================================
    # SLIDE 13: SOLUSI & VERIFIKASI FAULT 3
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13, COLOR_BG_LIGHT)
    add_header(s13, "Solusi & Verifikasi", "Perbaikan Defect 3: Compensating Transaction Rollback", 13)

    tf_sol_3, _ = add_card(s13, Inches(0.8), Inches(1.8), Inches(6.0), Inches(5.1), "Implementasi Kode Perbaikan (Fixed)", COLOR_SUCCESS_GREEN)
    code_fixed_3 = (
        "class FixedOrderService:\n"
        "    def checkout(self, customer, items, ...):\n"
        "        allocated_items = []\n"
        "        try:\n"
        "            # 1. Alokasi stok dengan pelacakan\n"
        "            for item in items:\n"
        "                if not self.inventory.deduct_stock(item.sku, item.qty):\n"
        "                    self._rollback(allocated_items)\n"
        "                    return order_failed\n"
        "                allocated_items.append(item)\n\n"
        "            # 2. Proses pembayaran\n"
        "            self.gateway.process_charge(...)\n"
        "            return order_success\n\n"
        "        except Exception as err:\n"
        "            # 3. KOMPENSASI OTOMATIS SAAT TERJADI ERROR\n"
        "            self._rollback(allocated_items)\n"
        "            order.status = OrderStatus.FAILED\n"
        "            return order"
    )
    p = tf_sol_3.add_paragraph()
    p.text = code_fixed_3
    p.font.name = FONT_CODE
    p.font.size = Pt(9)
    p.font.color.rgb = COLOR_SUCCESS_GREEN
    p.space_after = Pt(8)

    p = tf_sol_3.add_paragraph()
    p.text = "💡 Prinsip Solusi:\nPola Unit-of-Work / Saga Rollback. Setiap mutasi state yang telah dilakukan dicatat dan dijamin dipulihkan (*compensated*) secara otomatis jika langkah berikutnya mengalami kegagalan."
    p.font.name = FONT_FAMILY
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_MAIN

    tf_ver_3, _ = add_card(s13, Inches(7.1), Inches(1.8), Inches(5.2), Inches(5.1), "Hasil Verifikasi Otomatis (test_faults_fixed.py)", COLOR_ACCENT_BLUE)
    p = tf_ver_3.add_paragraph()
    p.text = "Hasil Eksekusi Unit Test Verifikasi:\n"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.bold = True

    verif_items_3 = [
        ("Stok Awal Gudang", "5 Unit"),
        ("Jumlah Pesanan", "2 Unit"),
        ("Simulasi Gateway", "HTTP 504 Gateway Timeout"),
        ("Status Pesanan", "FAILED (Dibatalkan Aman)"),
        ("Aksi Sistem", "Compensating Rollback Dieksekusi"),
        ("Stok Pasca Rollback", "5 Unit (Kembali Utuh 100%!)"),
        ("Phantom Deduction", "0 Unit (Tereliminasi Total)"),
        ("Status Pengujian", "PASSED / OK (Data Integrity Terjamin)")
    ]
    for k, v in verif_items_3:
        p = tf_ver_3.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True if "Status" in k or "Pasca" in k else False
        p.font.color.rgb = COLOR_SUCCESS_GREEN if "Status" in k or "Pasca" in k else COLOR_TEXT_MAIN
        p.space_after = Pt(4)

    # =========================================================================
    # SLIDE 14: MATRIKS CAPA (CORRECTIVE & PREVENTIVE ACTION)
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_background(s14, COLOR_BG_LIGHT)
    add_header(s14, "Action Plan Manajemen Mutu", "Matriks Tindakan Korektif & Preventif (CAPA)", 14)

    # Table layout for CAPA
    rows, cols = 4, 4
    left = Inches(0.8)
    top = Inches(1.8)
    width = Inches(11.733)
    height = Inches(5.0)

    table_shape = s14.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(2.7)
    table.columns[2].width = Inches(3.6)
    table.columns[3].width = Inches(3.633)

    headers_table = ["Defect ID", "Defect Description", "Corrective Action (Langsung)", "Preventive Action (Jangka Panjang)"]
    for c_idx, h_text in enumerate(headers_table):
        cell = table.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_WHITE

    capa_data = [
        ("DEF-01\n(Critical)", "Concurrency Race Condition (Overselling)", "Menerapkan threading.RLock pada modul inventory untuk melindungi operasi read-modify-write.", "Menambahkan Concurrency Stress Testing (Locust/k6) ke CI/CD; menggunakan Pessimistic Locking di Database."),
        ("DEF-02\n(High)", "Float Precision Leak & Unbounded Discount", "Refactoring ke decimal.Decimal dan membatasi total diskon dengan fungsi min(discount, subtotal).", "Mewajibkan Boundary Value Analysis (BVA) pada spesifikasi promosi; membuat Value Object Money terpusat."),
        ("DEF-03\n(Critical)", "Phantom Stock Deduction (No Rollback)", "Menambahkan blok kompensasi try-except-rollback pada service alur pemesanan checkout.", "Menerapkan arsitektur Saga Pattern untuk transaksi terdistribusi; menyusun cron job rekonsiliasi data stok.")
    ]

    for r_idx, (d_id, d_desc, d_ca, d_pa) in enumerate(capa_data, start=1):
        row_cells = [table.cell(r_idx, 0), table.cell(r_idx, 1), table.cell(r_idx, 2), table.cell(r_idx, 3)]
        contents = [d_id, d_desc, d_ca, d_pa]
        bg = COLOR_CARD_BG if r_idx % 2 == 1 else RGBColor(241, 245, 249)

        for c_idx, cell in enumerate(row_cells):
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            p = cell.text_frame.paragraphs[0]
            p.text = contents[c_idx]
            p.font.name = FONT_FAMILY
            p.font.size = Pt(10.5)
            p.font.color.rgb = COLOR_TEXT_MAIN
            if c_idx == 0:
                p.font.bold = True

    # =========================================================================
    # SLIDE 15: BEST PRACTICES PENGUJIAN PERANGKAT LUNAK
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_slide_background(s15, COLOR_BG_LIGHT)
    add_header(s15, "Rekomendasi Software Quality", "Best Practices Pengujian & Keandalan Sistem", 15)

    practices = [
        ("1. Pengujian Nilai Batas (Boundary Value Analysis)", "Mayoritas defect bisnis (seperti diskon melebihi 100% atau tagihan negatif) terjadi di titik batas ekstrem. Pengujian harus mencakup nilai: 0, nilai batas atas promo, angka negatif, dan nilai pecahan terkecil.", COLOR_ACCENT_BLUE),
        ("2. Concurrency & Multi-Thread Testing", "Jangan mengandalkan pengujian single-threaded sederhana. Gunakan mekanisme Barrier Synchronization atau Thread Pool untuk mensimulasikan ribuan transaksi simultan pada resource terbatas.", COLOR_ACCENT_BLUE),
        ("3. Chaos Engineering & Fault-Injection", "Uji ketahanan sistem terhadap skenario terburuk secara sengaja (koneksi payment putus di tengah jalan, database timeout, server restart). Pastikan sistem selalu kembali ke state konsisten.", COLOR_ACCENT_BLUE),
        ("4. Automated CI/CD Quality Gates", "Terapkan pipeline pengujian otomatis sebelum kode dapat di-merge: linter tipe data statis (mypy), automated regression test, dan mutation testing untuk mengukur efektivitas test suite.", COLOR_ACCENT_BLUE)
    ]

    for i, (p_title, p_desc, col) in enumerate(practices):
        col_idx = i % 2
        row_idx = i // 2
        left = Inches(0.8 + col_idx * 5.9)
        top = Inches(1.8 + row_idx * 2.5)
        tf, _ = add_card(s15, left, top, Inches(5.6), Inches(2.2), p_title, col)

        p = tf.add_paragraph()
        p.text = p_desc
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 16: KESIMPULAN & PENUTUP (Q&A)
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    set_slide_background(s16, COLOR_PRIMARY_DARK)

    # Accent decorative bar
    bar = s16.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.5), Inches(0.15), Inches(4.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_SUCCESS_GREEN
    bar.line.fill.background()

    # Closing content
    close_box = s16.shapes.add_textbox(Inches(1.4), Inches(1.5), Inches(10.5), Inches(4.5))
    tf16 = close_box.text_frame
    tf16.word_wrap = True

    p = tf16.paragraphs[0]
    p.text = "KESIMPULAN & LESSON LEARNED"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_CYAN
    p.space_after = Pt(12)

    conclusions = [
        "1. Kualitas perangkat lunak tidak hanya diukur dari kode yang bebas error sintaks, tetapi dari ketahanan logika terhadap konkurensi tinggi, integritas data moneter, dan penanganan kegagalan parsial.",
        "2. Metodologi Root Cause Analysis (5-Whys & Fishbone Diagram) terbukti efektif menembus akar masalah terdalam dari sekadar symptom fisik hingga proses pengembangan & arsitektur sistem.",
        "3. Melalui verifikasi automated testing, seluruh 3 defect berhasil diselesaikan 100% tanpa regresi.",
    ]
    for c in conclusions:
        p = tf16.add_paragraph()
        p.text = c
        p.font.name = FONT_FAMILY
        p.font.size = Pt(13)
        p.font.color.rgb = RGBColor(226, 232, 240)
        p.space_after = Pt(8)

    p = tf16.add_paragraph()
    p.text = "\nTERIMA KASIH — SESI TANYA JAWAB (Q & A)"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_WHITE
    p.space_after = Pt(6)

    p = tf16.add_paragraph()
    p.text = "Zainul Mutawakkil (NIM: 23EO10003)  •  Mata Kuliah: Analisis dan Pengujian Sistem"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(148, 163, 184)

    # Save presentation
    output_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "RCA_Presentasi_Tugas_AP.pptx"))
    prs.save(output_path)
    print(f"File presentasi berhasil dibuat: {output_path}")

if __name__ == "__main__":
    create_deck()
